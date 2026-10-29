"""把 Marp 簡報轉成 PPTX：每頁是一張圖片，<video> 頁會嵌入真正的影片（進入該頁自動播放）。

用法（在專案根目錄）：
    python3 tools/md2pptx.py decks/2026-10-30/2026-10-30.md

需要：Node.js（npx）、python-pptx（pip install python-pptx pillow）、Google Chrome。
影片請用 H.264 的 .mp4（iPhone 的 HEVC .mov 在 Windows PowerPoint 常常放不出來），
macOS 可用：avconvert -s IN.mov -p Preset1920x1080 -o OUT.mp4
"""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

MAGENTA = (255, 0, 255)
VIDEO_RE = re.compile(r'<video\b([^>]*)>\s*</video>', re.S)
ATTR_RE = re.compile(r'(\w+)="([^"]*)"')

AUTOPLAY_TIMING = """<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
<p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""


def magenta_bbox(img):
    """回傳洋紅色佔位框的 (left, top, right, bottom)，沒有就回傳 None。"""
    rgb = img.convert("RGB")
    mask = Image.eval(rgb.split()[0], lambda v: 255 if v > 245 else 0)
    g = Image.eval(rgb.split()[1], lambda v: 255 if v < 10 else 0)
    b = Image.eval(rgb.split()[2], lambda v: 255 if v > 245 else 0)
    from PIL import ImageChops
    mask = ImageChops.multiply(ImageChops.multiply(mask, g), b)
    return mask.getbbox()


def main(md_path, out_path=None):
    md = Path(md_path).resolve()
    out = Path(out_path) if out_path else md.with_suffix(".pptx")
    text = md.read_text(encoding="utf-8")

    # 影片換成同尺寸的洋紅色方塊，渲染後用來找影片位置
    videos = []

    def to_box(m):
        attrs = dict(ATTR_RE.findall(m.group(1)))
        videos.append(attrs)
        return f'<div style="{attrs.get("style", "")};background:#ff00ff"></div>'

    tmp_md = md.with_name(f"_pptx_{md.name}")
    tmp_md.write_text(VIDEO_RE.sub(to_box, text), encoding="utf-8")
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run(
                ["npx", "-y", "@marp-team/marp-cli@4.0.3", str(tmp_md), "--images", "png",
                 "--image-scale", "2", "--allow-local-files", "-o", f"{tmp}/s.png"],
                check=True, cwd=next(d for d in md.parents if (d / ".marprc.yml").exists()))
        finally:
            tmp_md.unlink()
        pngs = sorted(Path(tmp).glob("s.*.png"))

        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
        blank = prs.slide_layouts[6]
        vid = iter(videos)
        for png in pngs:
            slide = prs.slides.add_slide(blank)
            img = Image.open(png)
            sx = prs.slide_width / img.width
            sy = prs.slide_height / img.height
            slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
            box = magenta_bbox(img)
            if not box:
                continue
            v = next(vid)
            l, t, r, b = box
            movie = slide.shapes.add_movie(
                str(md.parent / v["src"]), int(l * sx), int(t * sy),
                int((r - l) * sx), int((b - t) * sy),
                poster_frame_image=str(md.parent / v["poster"]) if "poster" in v else None,
                mime_type="video/mp4")
            # 換成「進入此頁自動播放」的動畫時序
            old = slide.element.find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}timing")
            if old is not None:
                slide.element.remove(old)
            slide.element.append(etree.fromstring(AUTOPLAY_TIMING.format(spid=movie.shape_id)))
            print(f"{png.name}: 嵌入 {v['src']}")
        prs.save(out)
    print(f"完成：{out}")


if __name__ == "__main__":
    main(*sys.argv[1:])
