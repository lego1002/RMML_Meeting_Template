"""沒有 Node.js 時的備案：用已匯出的 PDF 當底圖做 PPTX，<video> 頁嵌入真正的影片（進入該頁自動播放）。

用法（在專案根目錄）：
    python3 tools/pdf2pptx.py decks/2026-10-30/2026-10-30.md

需要：python-pptx、pymupdf、pillow、lxml（pip install python-pptx pymupdf pillow lxml）。
PDF 要先用 marp-cli 匯出、且比 .md 新；影片位置取自 PDF 上 poster 截圖的位置。
"""
import sys
import tempfile
from pathlib import Path

import fitz
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).parent))
from md2pptx import ATTR_RE, AUTOPLAY_TIMING, VIDEO_RE  # noqa: E402


def slide_videos(text):
    """回傳 {頁碼(0 起算): video 屬性}。"""
    body = text.split("\n---\n", 1)[1] if text.startswith("---") else text
    out = {}
    for i, chunk in enumerate(body.split("\n---\n")):
        m = VIDEO_RE.search(chunk)
        if m:
            out[i] = dict(ATTR_RE.findall(m.group(1)))
    return out


def main(md_path, out_path=None):
    md = Path(md_path).resolve()
    pdf = md.with_suffix(".pdf")
    out = Path(out_path) if out_path else md.with_suffix(".pptx")
    if not pdf.exists() or pdf.stat().st_mtime < md.stat().st_mtime:
        sys.exit(f"{pdf.name} 不存在或比 .md 舊，請先重新匯出 PDF")
    videos = slide_videos(md.read_text(encoding="utf-8"))
    doc = fitz.open(pdf)
    if max(videos, default=0) >= len(doc):
        sys.exit("PDF 頁數跟 .md 對不上，請重新匯出 PDF")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory() as tmp:
        for i, page in enumerate(doc):
            png = Path(tmp) / f"s{i + 1:02d}.png"
            page.get_pixmap(dpi=144).save(png)  # 960×540 pt → 1920×1080
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
            if i not in videos:
                continue
            v = videos[i]
            # poster 是該頁最大的圖（排除整頁背景）
            pw, ph = page.rect.width, page.rect.height
            boxes = [im["bbox"] for im in page.get_image_info()
                     if (im["bbox"][2] - im["bbox"][0]) < pw * 0.95]
            l, t, r, b = max(boxes, key=lambda bb: (bb[2] - bb[0]) * (bb[3] - bb[1]))
            sx, sy = prs.slide_width / pw, prs.slide_height / ph
            movie = slide.shapes.add_movie(
                str(md.parent / v["src"]), int(l * sx), int(t * sy),
                int((r - l) * sx), int((b - t) * sy),
                poster_frame_image=str(md.parent / v["poster"]) if "poster" in v else None,
                mime_type="video/mp4")
            old = slide.element.find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}timing")
            if old is not None:
                slide.element.remove(old)
            slide.element.append(etree.fromstring(AUTOPLAY_TIMING.format(spid=movie.shape_id)))
            print(f"第 {i + 1} 頁：嵌入 {v['src']}")
    prs.save(out)
    print(f"完成：{out}")


if __name__ == "__main__":
    main(*sys.argv[1:])
